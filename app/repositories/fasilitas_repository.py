import os
from pathlib import Path

from openpyxl import load_workbook

from app.models.fasilitas import Fasilitas

EXCEL_PATH_ENV = "FASILITAS_EXCEL_PATH"
DEFAULT_EXCEL_RELATIVE_PATH = Path("Data") / "Data_Fasilitas_BSU_KLHK_2025.xlsx"
SHEET_NAME = "Data"
FACILITY_TYPE = "BSU"
FACILITY_STATUS = "A"
COLUMN_TAHUN = "tahun"
COLUMN_PROVINCE = "nama_propinsi"
COLUMN_REGENCY = "nama_dati2"
COLUMN_NAME = "nama"
COLUMN_TYPE = "jenis"
COLUMN_STATUS = "status"
COLUMN_SCOPE = "wilayah"
COLUMN_MANAGED_BY = "kelola"
COLUMN_ADDRESS = "alamat"
COLUMN_VILLAGE = "kelurahan"
COLUMN_DISTRICT = "kecamatan"
COLUMN_LATITUDE = "lat"
COLUMN_LONGITUDE = "lng"
COLUMN_WASTE_RECEIVED = "sampah_diterima_tahun"
COLUMN_WASTE_MANAGED = "jml_kelola_tahun"
FALLBACK_NAME = "Bank Sampah Tanpa Nama"
ADDRESS_SEPARATOR = ", "


class FasilitasRepository:
    def __init__(self, excel_path: Path | None = None) -> None:
        self._excel_path = excel_path or self._resolve_path()

    @property
    def excel_path(self) -> Path:
        return self._excel_path

    def load(self) -> list[Fasilitas]:
        if not self._excel_path.is_file():
            raise FileNotFoundError(f"File data fasilitas tidak ditemukan: {self._excel_path}")

        workbook = load_workbook(self._excel_path, read_only=True, data_only=True)
        try:
            sheet = workbook[SHEET_NAME]
            rows = sheet.iter_rows(values_only=True)
            header = next(rows, None)
            if header is None:
                raise ValueError(f"Sheet '{SHEET_NAME}' pada {self._excel_path} kosong.")

            columns = self._map_columns(header)
            facilities: dict[tuple[str, float, float], Fasilitas] = {}
            for excel_row_number, row in enumerate(rows, start=2):
                facility = self._build_fasilitas(row, columns, excel_row_number)
                if facility is not None:
                    facilities.setdefault(self._identity(facility), facility)
            return list(facilities.values())
        finally:
            workbook.close()

    def _resolve_path(self) -> Path:
        configured_path = os.getenv(EXCEL_PATH_ENV)
        if configured_path:
            return Path(configured_path).expanduser().resolve()
        return (Path(__file__).resolve().parents[2] / DEFAULT_EXCEL_RELATIVE_PATH).resolve()

    def _map_columns(self, header: tuple) -> dict[str, int]:
        normalized = [str(cell).strip().lower() if cell is not None else "" for cell in header]
        return {name: normalized.index(name) for name in normalized if name}

    def _build_fasilitas(
        self,
        row: tuple,
        columns: dict[str, int],
        excel_row_number: int,
    ) -> Fasilitas | None:
        latitude = self._to_float(self._value(row, columns, COLUMN_LATITUDE))
        longitude = self._to_float(self._value(row, columns, COLUMN_LONGITUDE))
        if latitude is None or longitude is None:
            return None

        facility_type = self._to_text(self._value(row, columns, COLUMN_TYPE))
        if facility_type and facility_type.upper() != FACILITY_TYPE:
            return None

        status = self._to_text(self._value(row, columns, COLUMN_STATUS))
        if status and status.upper() != FACILITY_STATUS:
            return None

        province = self._to_text(self._value(row, columns, COLUMN_PROVINCE))
        regency = self._to_text(self._value(row, columns, COLUMN_REGENCY))
        village = self._to_text(self._value(row, columns, COLUMN_VILLAGE))
        district = self._to_text(self._value(row, columns, COLUMN_DISTRICT))

        return Fasilitas(
            id=f"BSU-{excel_row_number:05d}",
            name=self._to_text(self._value(row, columns, COLUMN_NAME)) or FALLBACK_NAME,
            address=self._build_address(
                self._to_text(self._value(row, columns, COLUMN_ADDRESS)),
                village,
                district,
            ),
            province=province,
            regency=regency,
            village=village,
            district=district,
            latitude=latitude,
            longitude=longitude,
            managed_by=self._to_text(self._value(row, columns, COLUMN_MANAGED_BY)),
            scope=self._to_text(self._value(row, columns, COLUMN_SCOPE)),
            waste_received=self._to_float(self._value(row, columns, COLUMN_WASTE_RECEIVED)),
            waste_managed=self._to_float(self._value(row, columns, COLUMN_WASTE_MANAGED)),
            year=self._to_int(self._value(row, columns, COLUMN_TAHUN)) or 0,
        )

    def _identity(self, facility: Fasilitas) -> tuple[str, float, float]:
        return (facility.name.strip().upper(), facility.latitude, facility.longitude)

    def _build_address(self, address: str | None, village: str | None, district: str | None) -> str:
        parts = [address, village, district]
        unique_parts: list[str] = []
        for part in parts:
            cleaned = (part or "").strip()
            if cleaned and not any(cleaned.lower() in existing.lower() for existing in unique_parts):
                unique_parts.append(cleaned)
        return ADDRESS_SEPARATOR.join(unique_parts)

    def _value(self, row: tuple, columns: dict[str, int], column: str):
        index = columns.get(column)
        if index is None or index >= len(row):
            return None
        return row[index]

    def _to_text(self, value) -> str | None:
        if value is None:
            return None
        text = str(value).strip()
        return text or None

    def _to_float(self, value) -> float | None:
        if value is None or isinstance(value, bool):
            return None
        if isinstance(value, (int, float)):
            return float(value)
        try:
            return float(str(value).strip().replace(",", "."))
        except ValueError:
            return None

    def _to_int(self, value) -> int | None:
        number = self._to_float(value)
        return int(number) if number is not None else None